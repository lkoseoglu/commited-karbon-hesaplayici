#!/usr/bin/env python3
"""
HCMI v2.0 veri dosyasi uretici
==============================

Resmi "HCMI v2.0 Tool.xlsx" dosyasindaki SABIT verileri (emisyon faktorleri,
birim donusumleri, sogutucu GWP degerleri, camasirhane katsayilari, carpanlar)
okuyup tarayicida ve testlerde kullanilan tek bir JavaScript dosyasina yazar:

    assets/js/hcmi-data.js

KULLANIM
    python3 tools/build-hcmi-data.py --tool "/yol/HCMI v2.0 Tool.xlsx"
    python3 tools/build-hcmi-data.py --tool "/yol/HCMI v2.0 Tool.xlsx" --check   # dosyayi degistirmeden kontrol eder

NOT
  * Resmi Excel dosyasi depoda TUTULMAZ (telif + 5 MB). Sustainable Hospitality
    Alliance sitesinden indirilir; uretilen dosyanin icine dosyanin SHA-256
    degeri yazilir, boylece hangi surumden uretildigi izlenebilir.
  * Bu betik Excel'deki formullerden de sabitleri okur (ornegin camasirhane
    180 kWh/ton) ve beklenen degerle karsilastirir; resmi dosya degisirse
    betik sessizce yanlis veri uretmek yerine hata verir.
  * Uretilen dosyayi elle DUZENLEMEYIN; degisiklik gerekiyorsa bu betigi
    (veya assets/js/tr-reference.js icindeki Turkiye katmanini) degistirin.
"""
import argparse
import datetime as dt
import hashlib
import json
import os
import re
import sys
import warnings

warnings.filterwarnings("ignore")
import openpyxl  # noqa: E402

# --------------------------------------------------------------------------
# Elle girilen sozluk: enerji turu kimligi, Turkce ad, grup, varsayilan birim.
# Anahtar = Excel'deki resmi ad (degisirse betik hata verir).
# --------------------------------------------------------------------------
ENERGY_META = {
    "Purchased Electricity (Grid)":     ("grid",            "Şebekeden satın alınan elektrik",                       "electricity", "kWh"),
    "Natural Gas":                      ("natural_gas",     "Doğal gaz",                                             "gas",         "CuM"),
    "Butane":                           ("butane",          "Bütan",                                                 "gas",         "kg"),
    "Propane":                          ("propane",         "Propan",                                                "gas",         "kg"),
    "Liquefied Petroleum Gas (LPG)":    ("lpg",             "LPG (tüp / dökme)",                                     "gas",         "kg"),
    "Liquefied Natural Gas (LNG)":      ("lng",             "LNG (sıvılaştırılmış doğal gaz)",                       "gas",         "kg"),
    "Compressed Natural Gas (CNG)":     ("cng",             "CNG (sıkıştırılmış doğal gaz)",                         "gas",         "CuM"),
    "Gasoline (Stationary)":            ("gasoline_stat",   "Benzin (sabit tesis)",                                  "liquid",      "L"),
    "Diesel (Stationary)":              ("diesel_stat",     "Dizel / motorin (jeneratör, kazan vb.)",                "liquid",      "L"),
    "Fuel Oil 1":                       ("fuel_oil_1",      "Fuel oil 1",                                            "liquid",      "L"),
    "Fuel Oil 2":                       ("fuel_oil_2",      "Fuel oil 2 (kalorifer yakıtı / fuel-oil)",              "liquid",      "L"),
    "Fuel Oil 4":                       ("fuel_oil_4",      "Fuel oil 4",                                            "liquid",      "L"),
    "Fuel Oil 5":                       ("fuel_oil_5",      "Fuel oil 5",                                            "liquid",      "L"),
    "Fuel Oil 6":                       ("fuel_oil_6",      "Fuel oil 6 (ağır fuel-oil)",                            "liquid",      "L"),
    "Towngas (HK)":                     ("towngas_hk",      "Towngas (Hong Kong şehir gazı)",                        "gas",         "Unit"),
    "Tokyo Gas (Tokyo)":                ("tokyo_gas",       "Tokyo Gas (Tokyo şehir gazı)",                          "gas",         "CuM"),
    "Purchased Steam":                  ("steam",           "Satın alınan buhar",                                    "district",    "kWh"),
    "Purchased Chilled Water":          ("chilled_water",   "Satın alınan soğutulmuş su (bölgesel soğutma)",         "district",    "kWh"),
    "Purchased Hot Water":              ("hot_water",       "Satın alınan sıcak su (bölgesel ısıtma)",               "district",    "kWh"),
    "Biofuel Landfill gas (50/50)":     ("landfill_gas",    "Biyoyakıt: çöp depolama sahası gazı (50/50)",           "gas",         "CuM"),
    "Biogas (captured methane)":        ("biogas",          "Biyogaz (yakalanan metan)",                             "gas",         "CuM"),
    "Biofuel Used Oil":                 ("used_oil",        "Biyoyakıt: kullanılmış yağ",                            "liquid",      "L"),
    "Biofuel Wood Waste":               ("wood_waste",      "Biyoyakıt: ahşap atığı",                                "solid",       "kg"),
    "Biomass":                          ("biomass",         "Biyokütle (odun, pelet vb.)",                           "solid",       "kg"),
    "Gasoline (Mobile) Biofuel Blend":  ("gasoline_bio",    "Benzin (araç) – biyoyakıt karışımı",                    "mobile",      "L"),
    "Gasoline (Mobile)":                ("gasoline_mob",    "Benzin (araç / bahçe ekipmanı)",                        "mobile",      "L"),
    "Diesel (Mobile)":                  ("diesel_mob",      "Dizel / motorin (araç / iş makinesi)",                  "mobile",      "L"),
    "Charcoal":                         ("charcoal",        "Odun kömürü",                                           "solid",       "kg"),
    "Kerosene":                         ("kerosene",        "Gazyağı",                                               "liquid",      "L"),
    "Ethanol":                          ("ethanol",         "Etanol",                                                "liquid",      "L"),
    "Onsite Renewable Energy":          ("onsite_renew",    "Tesiste üretilen yenilenebilir enerji (GES vb.)",       "electricity", "kWh"),
    "Purchased Electricity (PPA)":      ("ppa",             "Elektrik – PPA (yenilenebilir alım anlaşması)",         "electricity", "kWh"),
}

MOBILE_NAMES = {"Gasoline (Mobile) Biofuel Blend", "Gasoline (Mobile)", "Diesel (Mobile)"}

UNIT_TR = {
    "DTH": "Dekatherm", "GJ": "Gigajul", "kBtu": "Bin Btu", "kWh": "Kilovatsaat (kWh)", "MJ": "Megajul",
    "MMBtu": "Milyon Btu", "MWh": "Megavatsaat (MWh)", "Therm": "Therm", "Ton-Hr": "Ton-saat (soğutma)",
    "Unit": "Birim (Towngas)", "MCF": "Bin fit³", "CCF": "Yüz fit³", "cuft": "Fit³", "cuyd": "Yarda³",
    "CuM": "Metreküp (m³)", "Bar": "Varil", "ML": "Megalitre", "kL": "Kilolitre (m³ sıvı)", "L": "Litre",
    "Gal": "Galon (ABD)", "kGal": "Bin galon", "kg": "Kilogram", "klb": "Bin pound", "lb": "Pound",
    "MT": "Metrik ton", "UK ton": "İngiliz tonu", "US ton": "Amerikan tonu",
}

# Excel'den okunacak ve beklenen sabitler (dogrulama icin)
EXPECT = {
    "laundry_ton_per_room": 0.00512559378,
    "us_ton_to_metric": 0.90718474,
    "laundry_kwh_elec_per_t": 180,
    "laundry_kwh_gas_per_t": 1560,
    "laundry_l_fueloil2_per_t": 111,
    "uplift_both": 1.02,
    "uplift_one": 1.01,
    "meeting_days": 365,
    "meeting_hours": 10,
}


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def num(v):
    """Excel hucresini JS'e guvenli sayiya cevirir: sayi -> float, yoksa None."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if not float(v).is_integer() else (int(v) if abs(v) < 1e15 else float(v))
    return None


def js(value):
    """Python -> JS literal (deterministik, en kisa gidis-donus)."""
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), sort_keys=False)


def must_find(pattern, text, what):
    m = re.search(pattern, text)
    if not m:
        sys.exit(f"HATA: '{what}' formulde bulunamadi (desen: {pattern}). Resmi dosya degismis olabilir.")
    return m


def build(tool_path):
    wb_v = openpyxl.load_workbook(tool_path, data_only=True)
    wb_f = openpyxl.load_workbook(tool_path, data_only=False)

    # ---------------------------------------------------------------- birimler
    uc = wb_v["UnitConversions"]
    type_names = [uc.cell(row=2, column=c).value for c in range(4, 36)]  # D..AI
    if len(type_names) != 32 or any(n not in ENERGY_META for n in type_names):
        missing = [n for n in type_names if n not in ENERGY_META]
        sys.exit(f"HATA: enerji turu listesi beklenenden farkli. Eksik sozluk girdisi: {missing}")

    units = []
    for r in range(4, 31):
        cat, name, abb = uc.cell(row=r, column=1).value, uc.cell(row=r, column=2).value, uc.cell(row=r, column=3).value
        if not abb:
            continue
        units.append({"abb": abb, "name": name, "tr": UNIT_TR.get(abb, name), "cat": cat})
    if len(units) != 27:
        sys.exit(f"HATA: 27 birim bekleniyordu, {len(units)} bulundu.")

    energy_types = []
    for ci, name in enumerate(type_names):
        col = 4 + ci
        conv = {}
        for r in range(4, 31):
            abb = uc.cell(row=r, column=3).value
            v = num(uc.cell(row=r, column=col).value)
            if abb and v is not None:
                conv[abb] = v
        eid, tr, grp, dunit = ENERGY_META[name]
        if dunit not in conv:
            sys.exit(f"HATA: '{name}' icin varsayilan birim '{dunit}' donusum tablosunda yok.")
        energy_types.append({
            "id": eid, "name": name, "tr": tr, "group": grp,
            "stationary": name not in MOBILE_NAMES,
            "defaultUnit": dunit,
            "kwhPerUnit": conv,
        })

    # --------------------------------------------------------------- EF tablosu
    # Not: EFs sayfasinin sutun sirasi UnitConversions ile ayni DEGIL ("Gasoline (Mobile) Biofuel
    # Blend" farkli yerde), bu yuzden sutunlar ad ile eslestirilir.
    ef = wb_v["EFs"]
    hdr = {ef.cell(row=4, column=c).value: c for c in range(3, 35)}
    if ef.cell(row=4, column=1).value != "Country" or ef.cell(row=4, column=2).value != "Province/Region":
        sys.exit("HATA: EFs basliklari beklenen yerde degil.")
    if set(hdr) != set(type_names):
        sys.exit(f"HATA: EFs ile UnitConversions enerji turu adlari farkli: {set(hdr) ^ set(type_names)}")
    rows = []
    for r in range(6, ef.max_row + 1):
        c = ef.cell(row=r, column=1).value
        if c is None:
            continue
        rows.append((c, ef.cell(row=r, column=2).value, {n: ef.cell(row=r, column=col).value for n, col in hdr.items()}))
    if len(rows) != 337:
        sys.exit(f"HATA: 337 bolge bekleniyordu, {len(rows)} bulundu.")

    GRID, CHILLED, PPA = "Purchased Electricity (Grid)", "Purchased Chilled Water", "Purchased Electricity (PPA)"
    profile_names = [n for n in type_names if n not in (GRID, CHILLED, PPA)]
    profile_types = [ENERGY_META[n][0] for n in profile_names]

    profiles, profile_index, regions = [], {}, []
    for country, region, vals in rows:
        prof = tuple(num(vals[n]) for n in profile_names)
        if prof not in profile_index:
            profile_index[prof] = len(profiles)
            profiles.append(list(prof))
        grid, ppa = num(vals[GRID]), num(vals[PPA])
        reg = {"c": country, "r": region, "p": profile_index[prof], "g": grid, "ch": num(vals[CHILLED])}
        if ppa != grid:
            reg["ppa"] = ppa  # yalniz farkli ise (Singapur: PPA faktoru yok)
        regions.append(reg)
    keys = {(r["c"], r["r"]) for r in regions}
    if len(keys) != len(regions):
        sys.exit("HATA: (ulke, bolge) cifti benzersiz degil.")

    # ----------------------------------------------------------- sogutucu GWP
    tb_v = wb_v["Tab B - Refrigerants"]
    refrigerants = []
    for r in list(range(19, 23)) + list(range(25, 62)):
        name, gwp = tb_v.cell(row=r, column=7).value, tb_v.cell(row=r, column=8).value
        gwp = num(float(gwp)) if gwp not in (None, "") else None
        if not name or gwp is None:
            sys.exit(f"HATA: Tab B satir {r} okunamadi.")
        refrigerants.append({"id": f"rf{len(refrigerants) + 1:02d}", "name": name, "gwp": gwp, "common": r <= 22})
    if len(refrigerants) != 41:
        sys.exit(f"HATA: 41 sogutucu bekleniyordu, {len(refrigerants)} bulundu.")

    # --------------------------------------------- formullerden sabitler (dogrulama)
    ta = wb_f["Tab A - Outsourced Laundry "]
    h28 = str(ta["H28"].value)
    h33, h34, h35 = str(ta["H33"].value), str(ta["H34"].value), str(ta["H35"].value)
    j17 = str(ta["J17"].value)
    s1 = wb_f["1. Hotel Details and Result"]
    h47, h51 = str(s1["H47"].value), str(s1["H51"].value)

    consts = {
        "laundry_ton_per_room": float(must_find(r"\*\s*([0-9.]+)\s*,\s*\"\"\)", h28, "H28 ton/oda").group(1)),
        "us_ton_to_metric": float(must_find(r"H17\*\s*([0-9.]+)", j17, "J17 US ton").group(1)),
        "laundry_kwh_elec_per_t": float(must_find(r"\(\s*([0-9.]+)\s*\*\s*\$H\$28", h33, "H33").group(1)),
        "laundry_kwh_gas_per_t": float(must_find(r"\(\s*([0-9.]+)\s*\*\s*\$H\$28", h34, "H34").group(1)),
        "laundry_l_fueloil2_per_t": float(must_find(r"\(\s*([0-9.]+)\s*\*\s*\$H\$28", h35, "H35").group(1)),
        "uplift_both": float(must_find(r"\)\*([0-9.]+),IF\(OR", h47, "H47 %2 yukleme").group(1)),
        "uplift_one": float(must_find(r"\)\*([0-9.]+)\)\+'Tab B", h47, "H47 %1 yukleme").group(1)),
        "meeting_days": float(must_find(r"/([0-9]+)\*1000/([0-9]+)/", h51, "H51 gun").group(1)),
        "meeting_hours": float(must_find(r"/([0-9]+)\*1000/([0-9]+)/", h51, "H51 saat").group(2)),
    }
    for k, v in EXPECT.items():
        if abs(consts[k] - v) > 1e-12:
            sys.exit(f"HATA: sabit '{k}' beklenen {v}, Excel'de {consts[k]}.")

    # Yarim-tahmin camasirhane tablosunun hangi enerji turlerine bakildigini dogrula
    for cell, typ in (("G33", "Purchased Electricity (Grid)"), ("G34", "Natural Gas"), ("G35", "Fuel Oil 2")):
        if typ not in str(ta[cell].value):
            sys.exit(f"HATA: Tab A {cell} '{typ}' icermiyor.")

    # Sabit (statik) enerji turleri listesi (BackOfHouse H sutunu) ile kendi bayragimiz tutarli mi?
    boh = wb_v["BackOfHouse"]
    stationary_list = [boh.cell(row=r, column=8).value for r in range(3, 32)]
    ours = [t["name"] for t in energy_types if t["stationary"]]
    if sorted(stationary_list) != sorted(ours):
        sys.exit("HATA: BackOfHouse 'EnergyTypesStationary' listesi ile kendi 'stationary' bayragimiz uyusmuyor.")

    # Toggle metinleri (veri dogrulama)
    toggles = ["Yes (manually add)", "Yes (default to 1% of total emissions)", "No don't add"]
    dv_texts = []
    for dv in wb_f["1. Hotel Details and Result"].data_validations.dataValidation:
        if str(dv.sqref).startswith("H35"):
            dv_texts.append(dv.formula1)
    if not dv_texts or toggles[0] not in dv_texts[0] or toggles[2] not in dv_texts[0]:
        sys.exit("HATA: H35 acilir liste secenekleri beklenenden farkli.")

    data = {
        "meta": {
            "tool": "HCMI v2.0 Tool.xlsx",
            "toolVersion": "2.0",
            "efReportingYear": 2021,
            "publisher": "Sustainable Hospitality Alliance / WTTC (araç: Greenview)",
            "sha256": sha256(tool_path),
            "builtAt": dt.date.today().isoformat(),
            "builder": "tools/build-hcmi-data.py",
            "note": "Bu dosya otomatik uretilir; elle duzenlemeyin.",
        },
        "units": units,
        "energyTypes": energy_types,
        "profileTypes": profile_types,
        "profiles": profiles,
        "regions": regions,
        "refrigerants": refrigerants,
        "laundry": {
            "tonnePerOccupiedRoom": consts["laundry_ton_per_room"],
            "usTonToMetricTonne": consts["us_ton_to_metric"],
            "estimate": [
                {"type": "Purchased Electricity (Grid)", "perTonne": consts["laundry_kwh_elec_per_t"], "unit": "kWh"},
                {"type": "Natural Gas", "perTonne": consts["laundry_kwh_gas_per_t"], "unit": "kWh"},
                {"type": "Fuel Oil 2", "perTonne": consts["laundry_l_fueloil2_per_t"], "unit": "L"},
            ],
        },
        "uplift": {"one": consts["uplift_one"], "both": consts["uplift_both"]},
        "meeting": {"daysPerYear": int(consts["meeting_days"]), "hoursPerDay": int(consts["meeting_hours"])},
        "toggles": {"manual": toggles[0], "default1pct": toggles[1], "none": toggles[2]},
        "limits": {"energyRows": 7, "marketPurchaseRows": 2, "laundryRows": 3},
    }
    return data


def render(data):
    parts = []
    parts.append("/* =========================================================================\n"
                 "   HCMI v2.0 — resmi araçtan türetilmiş sabit veriler (OTOMATİK ÜRETİLDİ)\n"
                 "   Kaynak : HCMI v2.0 Tool.xlsx (Sustainable Hospitality Alliance / WTTC)\n"
                 f"   SHA-256: {data['meta']['sha256']}\n"
                 "   Üreten : tools/build-hcmi-data.py   — bu dosyayı elle düzenlemeyin.\n"
                 "   ========================================================================= */\n")
    parts.append("(function (root) {\n  'use strict';\n  root.HCMI_DATA = {\n")
    keys = list(data.keys())
    for i, k in enumerate(keys):
        v = data[k]
        if k in ("regions", "energyTypes", "units", "refrigerants", "profiles"):
            sep = ",\n"
            body = sep.join("    " + js(x) for x in v)
            parts.append(f"  {k}: [\n{body}\n  ]")
        else:
            parts.append(f"  {k}: {js(v)}")
        parts.append(",\n" if i < len(keys) - 1 else "\n")
    parts.append("  };\n})(typeof globalThis !== 'undefined' ? globalThis : self);\n")
    return "".join(parts)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--tool", required=True, help="Resmi 'HCMI v2.0 Tool.xlsx' dosyasının yolu")
    ap.add_argument("--out", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "assets", "js", "hcmi-data.js"))
    ap.add_argument("--check", action="store_true", help="Dosyayı yazmadan, mevcut çıktıyla aynı olup olmadığını denetle")
    args = ap.parse_args()

    data = build(args.tool)
    text = render(data)
    out = os.path.abspath(args.out)
    if args.check:
        cur = open(out, encoding="utf-8").read() if os.path.exists(out) else ""
        strip = lambda s: re.sub(r"builtAt\":\"[0-9-]+\"", "builtAt", s)
        same = strip(cur) == strip(text)
        print("AYNI" if same else "FARKLI", out)
        sys.exit(0 if same else 1)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, "w", encoding="utf-8") as f:
        f.write(text)
    print(f"yazildi: {out}  ({len(text)/1024:.1f} KB)  bolge={len(data['regions'])} profil={len(data['profiles'])} "
          f"enerji={len(data['energyTypes'])} birim={len(data['units'])} sogutucu={len(data['refrigerants'])}")


if __name__ == "__main__":
    main()
